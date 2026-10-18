#!/usr/bin/env python3
"""Validate/render the authored ace specification. Never modifies game files."""
import argparse
import collections
import csv
import hashlib
import json
from pathlib import Path

ROLES = {'gym': '관장', 'elite': '사천왕', 'champion': '챔피언'}
PHASES = {'guard': '전술 방어', 'recharge': '재충전 휴식', 'sleep-rest': '자기 수면 휴식',
          'charge-electro': '일렉트로빔 충전(비면 생략)', 'charge-solar': '솔라빔 충전(쾌청이면 생략)',
          'charge-sky': '불새 충전'}
KINDS = {'attack': '공격', 'charge': '강공격', 'barrage': '연타', 'block': '전술 방어',
         'boost': '위력 보정', 'fortify': '방벽', 'recover': '회복', 'status': '상태 지시'}
STYLE_NAMES = {'tempo': '연타형', 'bulwark': '철벽형', 'precision': '특방형', 'sustain': '회복형',
               'scorch': '화상형', 'venom': '독형', 'roots': '씨앗형', 'focus': '집중형', 'assault': '돌파형'}


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def validate(game, spec, refs, csv_dir=None):
    dex = read(game/'data/pokedex_expansion.json')
    pokemon = {p['id']: p for p in dex['pokemon']}
    moves = dex['move_definitions']
    parties = read(game/'data/boss_parties.json')['entries']
    entries = spec['entries']
    references = {r['id']: r for r in refs['entries']}
    assert spec['status'] == 'specification_not_applied'
    assert spec['runtime_route_version'] == 5
    assert len(entries) == 137 and len({e['key'] for e in entries}) == 137
    assert {e['key'] for e in entries} == set(parties)
    assert len(references) == len(refs['entries']) == 138
    assert collections.Counter(e['role'] for e in entries) == {'gym': 86, 'elite': 36, 'champion': 15}
    legal_rows = set()
    all_slots = []
    used_references = set()
    for e in entries:
        assert e['key'].split(':')[0] == e['role']
        assert [a['species'] for a in e['aces']] == parties[e['key']]['fixed_aces'], e['key']
        for a in e['aces']:
            species = a['species']
            assert a['theme'] and a['adaptation'] and species in pokemon
            assert len(set(a['moves'])) == len(a['moves']) and 1 <= len(a['moves']) <= 4
            r = references[a['reference']]
            used_references.add(a['reference'])
            assert (r['key'], r['species']) == (e['key'], species)
            assert r['source'].startswith('https://') and r['source_game'] and r['source_battle']
            assert 1 <= len(r['source_moves']) <= 4
            assert all(str(i) in moves for i in r['source_moves'])
            assert 3 <= len(a['cycle']) <= 8
            assert all(p['phase'] in set(PHASES)|{'move'} for p in a['cycle'])
            assert all(p.get('move') in a['moves'] for p in a['cycle'] if p['phase'] == 'move')
            assert set(a['moves']) <= {p.get('move') for p in a['cycle']}, e['key']
            assert [v['move'] for v in a['move_audit']] == a['moves']
            for v in a['move_audit']:
                move_id = v['move']
                assert str(move_id) in moves
                assert v['current_selected_learnset'] == (move_id in dex['progression'][str(species)]['compatible_moves'])
                assert set(v['required_mechanics']) <= set(spec['mechanics'])
                proof = v['learnset_evidence']
                path = v['inheritance_path']
                assert path[-1] == species and proof['pokemon_id'] == pokemon[path[0]]['source_pokemon_id']
                assert all(pokemon[b]['evolves_from']==a for a,b in zip(path,path[1:]))
                assert proof['move_id'] == move_id and proof['version_group_id'] > 0
                assert 1 <= proof['pokemon_move_method_id'] <= 11
                legal_rows.add(tuple((k, proof[k]) for k in sorted(proof)))
                assert str(move_id) in spec['player_card_audit']
                all_slots.append((e, a, v))
            current = a['current_v5_at_level100']
            assert current['level'] == 100 and current['style'] in STYLE_NAMES
            assert all(v['kind'] in KINDS and str(v['move_id']) in moves for v in current['intents'])
    assert used_references == set(references)
    assert len(all_slots) == 539
    # Boss route constraints and especially same-person role separation must remain explicit.
    by_key = {e['key']:e for e in entries}
    for pair in [('gym:blue','champion:blue'), ('gym:hala','elite:hala'), ('gym:olivia','elite:olivia')]:
        assert by_key[pair[0]]['aces'][0]['moves'] != by_key[pair[1]]['aces'][0]['moves']
    c = by_key['champion:cynthia']['aces'][0]
    assert c['moves'] == [407,89,53,416] and c['current_v5_at_level100']['style'] == 'sustain'
    assert all(p['phase']!='move' or p['move']!=105 for p in c['cycle'])
    # Check inherited moves against the evolution chain instead of rejecting valid egg moves.
    for key, move_id, path in [('elite:koga',98,[41,42,169]),('elite:drake',407,[371,372,373]),('champion:red',344,[172,25,26])]:
        ace=by_key[key]['aces'][0]
        assert move_id in ace['moves']
        assert next(v for v in ace['move_audit'] if v['move']==move_id)['inheritance_path']==path
    if csv_dir:
        for name, expected in spec['learnset_source']['files'].items():
            assert hashlib.sha256((csv_dir/name).read_bytes()).hexdigest() == expected, name
        found = set()
        wanted = {(dict(t)['pokemon_id'], dict(t)['move_id']) for t in legal_rows}
        for row in csv.DictReader((csv_dir/'pokemon_moves.csv').open(encoding='utf-8')):
            if (int(row['pokemon_id']),int(row['move_id'])) not in wanted:
                continue
            actual = {k:int(v) for k,v in row.items() if v}
            for proof in (t for t in legal_rows if dict(t)['pokemon_id']==actual['pokemon_id'] and dict(t)['move_id']==actual['move_id']):
                if all(actual.get(k)==v for k,v in proof):found.add(proof)
        assert found == legal_rows, 'Missing exact version/method/level learnset evidence'
    return pokemon, moves, references, all_slots, len(legal_rows)


def render(game, materials, csv_dir=None, workbook=None):
    spec = read(materials/'docs/boss-ace-tactics.json')
    refs = read(materials/'docs/boss-ace-references.json')
    pokemon, moves, references, slots, legal_pairs = validate(game,spec,refs,csv_dir)
    names = lambda values: ' · '.join(moves[str(i)]['name'] for i in values)
    def cycle(a):
        return ' → '.join(moves[str(v['move'])]['name'] if v['phase']=='move' else PHASES[v['phase']] for v in a['cycle'])
    def gaps(a):
        tags = {t for v in a['move_audit'] for t in v['required_mechanics']}
        return ' / '.join('P%d %s' % (spec['mechanics'][t]['priority'],spec['mechanics'][t]['name'])
                          for t in sorted(tags,key=lambda t:(spec['mechanics'][t]['priority'],t))) or '직접 피해 중심'
    def current(a):
        v = a['current_v5_at_level100']
        return STYLE_NAMES[v['style']]+' · '+' → '.join(KINDS[p['kind']]+'('+moves[str(p['move_id'])]['name']+')' for p in v['intents'])
    entries = spec['entries']
    rollout_path = game/'data/boss_move_effects_v7.json'
    if not rollout_path.exists(): rollout_path = game/'data/boss_move_effects_v6.json'
    route_version = read(rollout_path)['route_version'] if rollout_path.exists() else 0
    deployed = set(read(rollout_path)['profiles']) if rollout_path.exists() else set()
    used = sorted({v['move'] for _,_,v in slots})
    historical = [(e,a,v) for e,a,v in slots if not v['current_selected_learnset']]
    changed = [(e,a) for e in entries for a in e['aces'] if a['moves'] != references[a['reference']]['source_moves']]
    lines = ['# 대표 포켓몬 기술·전술 명세', '',
        '**상태: 계획 1번의 조사·명세 완료. 게임 미적용.** main `'+spec['game_commit']+'`의 v5 전술은 그대로입니다. '
        '관장86·사천왕36·챔피언15 =137대진, 풍·란 두 대표를 포함해138마리입니다.', '',
        '원작 기술과 카드 전투용 행동 순서를 구분했습니다. 원작 기술은 승인 파티와 같은 작품/전투를 우선하며, '
        '행동 순서는 이 게임을 위한 고정 예고안입니다. 원작 AI가 이 순서로 움직였다는 뜻은 아닙니다. 원작이2~3기술만 쓰는 경우 빈 슬롯을 임의로 채우지 않았습니다.', '',
        '[검토 엑셀](boss-ace-tactics.xlsx) · [전술 JSON](boss-ace-tactics.json) · [원작 기술·고정 출처](boss-ace-references.json) · [승인 파티](boss-party-review.md)', '',
        '## 먼저 고칠 차이', '',
        '- 난천의 한카리아스는 현재 마지막 자리의 회복형을 물려받아 HP회복 행동을 합니다. 명세는 지진·드래곤다이브·화염방사·기가임팩트의 공격형입니다.',
        '- 꼭두는 승인 파티가 HGSS 재대결이므로 누르기·헤롱헤롱·잠꼬대·잠자기 기준입니다. 첫 체육관전의 구르기·우유마시기를 같은 전투의 기술처럼 섞지 않습니다.',
        '- 관장/챔피언 그린, 관장/사천왕 할라·라이치, 사천왕/챔피언 목호는 역할 키별로 기술과 전술을 따로 보관합니다.',
        '- 풍·란은 솔록→루나톤의 순차 출전입니다. 더블배틀을 도입하지 않습니다.', '',
        '## 각색 기준과 예외', '',
        '- 민화의 라플레시아·강연의 윈디는 승인 기준인 HGSS 재대결 팀에 없어 같은 인물의 레츠고 재대결 기술을 참조했습니다. 코르니의 루카리오는 체육관 밖 마스터타워 이벤트 기준이며 메가진화는 제외합니다.',
        '- 자포코일 기술은 원작 최종 레어코일을 기준으로 잡았습니다. 기존 파티 검토에 코일→자포코일로 표시한 진화 설명과 구분하며 승인된 실제 파티는 바꾸지 않았습니다.',
        '- 최종형으로 고정한 대표는 원작 미진화체 기술을 아래에 남기고 합법 습득이 확인된 각색 기술을 제안했습니다. 레드 라이츄는 원작 피카츄 기술을 계승합니다.',
        '- 크로뱃의 전광석화는 주뱃 교배→골뱃→크로뱃, 보만다의 드래곤다이브는 아공이 교배→쉘곤→보만다 경로로 계승할 수 있어 원작 기술을 유지합니다. 라이츄 볼트태클은 피츄 특수 교배→피카츄→라이츄, 모단단게 집게해머는 오기지게 레벨 습득→진화 경로도 남겼습니다.',
        f'- 제안539기술 슬롯({len(used)}종), 서로 다른 종/기술510쌍의 습득 기록(직접535슬롯·진화 계승4슬롯)을 고정 PokeAPI CSV에서 확인했습니다. 현재 종별 선택 버전에는 없는 과거 기술이{len(historical)}슬롯입니다. 보스의 역사적 기술 참조와 플레이어의 현행 훈련 가능 여부는 별개입니다. SV까지의 습득 근거이며 현재 버전의4기술 동시 공존·전송 제약까지 검증한 목록은 아닙니다.',
        '- 전술 방어는 공용 카드 전투 행동이며 리플렉터를 배운 것으로 표시하지 않습니다. 선공기라도 기존 아군턴/적턴 순서를 우회하지 않습니다.',
        '- 원작 확률·명중/회피·급소·PP를 그대로 도입하지 않습니다. 잠꼬대·변덕레이저 같은 선택/변동은 행동 예고에서 공개하도록 각색합니다. 구체 수치는 효과 구현 때 결정하며 이번에는 조정하지 않았습니다.',
        '- 일격필살은 공개 강공격으로 각색합니다. 테라버스트는 테라스탈이 없어 노말 특수 기술로 취급합니다. 메가/거다이맥스/테라스탈·특성·IV/EV·성격 보정은 추가하지 않습니다.',
        '- 마지막 대표의 볼트체인지/유턴은 피해만 남기는 각색안입니다. 상대 파티 순서와 대표 마지막 규칙을 깨는 자발 재출전은 넣지 않습니다. 강제 교대·설치기는 아군3마리/현재 교대 규칙을 기준으로 후속 구현합니다.', '',
        '## 효과 구현 순서(계획 2번의 입력)', '',
        '아래는 **미구현 또는 적 행동 연결이 필요한 범주**입니다. 플레이어 카드에 효과 키가 있어도 적 공격은 현재 직접 피해만 처리하므로 완료로 표시하지 않습니다. '
        '원작 효과 전체 재현을 약속하는 목록은 아니며, 각색이 필요한 항목도 함께 표시했습니다.', '',
        '| 순서 | 범주 | 필요한 동작 | 관련 대표 수 |', '|---|---|---|---|']
    for tag,m in sorted(spec['mechanics'].items(),key=lambda pair:(pair[1]['priority'],pair[0])):
        count = sum(any(tag in v['required_mechanics'] for v in a['move_audit']) for e in entries for a in e['aces'])
        lines.append(f"| P{m['priority']} | {m['name']} | {m['detail']} | {count} |")
    lines += ['', 'P0 기술별 적 효과 연결 → P1 회복/수면·충전/재충전·연타/반동 → P2 조건부 위력·교대/설치·보호 → P3 개별 각색 순으로 작업합니다. '
              '구현 시 새 전술 버전을 별도로 두고 v1~v5 저장의 원래 전술·RNG를 보존해야 합니다. 현재 적 정의가 저장된 전투뿐 아니라 다음 전투를 재생성하는 경우도 포함합니다.', '']
    for role in ROLES:
        lines += ['## '+ROLES[role], '']
        for region in dict.fromkeys(e['region'] for e in entries if e['role']==role):
            lines += ['### '+region, '', '| 대진·대표 | 원작 기술·참조 전투 | 제안 행동 순서 | 전술·차이 |', '|---|---|---|---|']
            for e in entries:
                if e['role']!=role or e['region']!=region:continue
                for a in e['aces']:
                    r=references[a['reference']]
                    ref='['+r['source_game']+' / '+r['source_battle']+']('+r['source']+')<br>'+names(r['source_moves'])
                    note=a['theme']+'<br>'+a['adaptation']
                    if r['basis']!='승인 파티와 같은 원작 전투':note+='<br>'+r['basis']
                    lines.append('| '+' | '.join([e['name']+' · **'+pokemon[a['species']]['name']+'**',ref,cycle(a),note])+' |')
            lines.append('')
    if deployed:
        lines[2] = f'**상태: 전체 목표 명세이며 일부 적용.** 현재 새 원정 v{route_version}에 대표{len(deployed)}마리의 단계별 기술 효과를 적용했습니다. [실제 적용 현황](boss-move-effects.md)을 확인하세요. 원작 참조와 전체 제안은 아래에 보존합니다. 나머지{138-len(deployed)}대표 전술은 이전 방식입니다.'
    lines += ['## 명세 작성 당시 검증과 한계', '',
              '- 137역할/138대표가 main 승인 파티의 고정 대표와 정확히 일치합니다. 제안 기술은1~4개, 행동 순서는 해당 대표의 제안 기술만 참조합니다.',
              f'- Godot4.7.2에서 현재 대표138마리의 Lv.18/50/100 전술414개가 기존 저장 검증 규칙에 유효함을 확인했습니다. 제안 기술{len(used)}종도 실제 카드 정의로 읽었습니다. 이는 새 전술을 게임에서 실행한 검사가 아닙니다.',
              '- 원작 참조의 종/폼·기술·출처와 제안의 습득 증거(버전/방법/레벨/진화 계승)를 JSON에 남겼습니다. 외부 원문/전체 데이터는 캐시이며 Git에 넣지 않습니다.',
              '- 계획1의 명세 작성 당시에는 런타임을 바꾸지 않았습니다. 후속 단계별 구현/실제 화면 검사는 [적용 현황](boss-move-effects.md)에 구분합니다. Windows/빌드/export/EXE 검사는 하지 않았습니다.', '']
    (materials/'docs/boss-ace-tactics.md').write_text('\n'.join(lines),encoding='utf-8')
    if workbook:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='대표 전술안'
        ws.append(['지방','역할','트레이너','대표','원작 기준','원작 기술','제안 기술','제안 행동 순서','전술','각색 이유','현재 Lv100 전술(미변경)','필요한 효과','출처','적용 현황'])
        for e in entries:
            for a in e['aces']:
                r=references[a['reference']]
                ws.append([e['region'],ROLES[e['role']],e['name'],pokemon[a['species']]['name'],r['source_game']+' / '+r['source_battle'],names(r['source_moves']),names(a['moves']),cycle(a),a['theme'],a['adaptation']+' / '+r['basis'],current(a),gaps(a),r['source'],f'v{route_version} 적용' if a['reference'] in deployed else '후속 구현'])
        ws=wb.create_sheet('기술 효과');ws.append(['ID','기술','타입','분류','원본 위력','현재 플레이어 카드 설명','적 기본 피해만 가능','후속 효과·각색','관련 대표 수'])
        for i in used:
            raw=moves[str(i)];tags={t for _,_,v in slots if v['move']==i for t in v['required_mechanics']}
            ws.append([i,raw['name'],raw['type'],{1:'변화',2:'물리',3:'특수'}[raw['category']],raw['power'],spec['player_card_audit'][str(i)]['description'],'가능' if raw['category']!=1 and raw['power']>0 else '별도 처리 필요',' / '.join(spec['mechanics'][t]['name'] for t in sorted(tags)),sum(i in a['moves'] for e in entries for a in e['aces'])])
        ws=wb.create_sheet('구현 순서');ws.append(['순서','범주','동작'])
        for tag,m in sorted(spec['mechanics'].items(),key=lambda pair:(pair[1]['priority'],pair[0])):ws.append(['P'+str(m['priority']),m['name'],m['detail']])
        for ws in wb:
            ws.freeze_panes='D2' if ws.title=='대표 전술안' else 'B2'
            ws.auto_filter.ref=ws.dimensions
            ws.sheet_view.zoomScale=85
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='314568')
            for row in ws.iter_rows(min_row=2):
                for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
            widths=[14,12,14,28,38,52,52,72,65,75,70,70,80,22] if ws.title=='대표 전술안' else ([9,24,14,12,12,65,24,75,15] if ws.title=='기술 효과' else [10,45,110])
            for col,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(col)].width=width
        wb.save(workbook)
    result={'roles':len(entries),'aces':138,'move_slots':len(slots),'unique_moves':len(used),'legal_pairs':legal_pairs,'historical_slots':len(historical),'adapted_aces':len(changed),'csv_revalidated':bool(csv_dir)}
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--game-dir',type=Path,required=True)
    p.add_argument('--csv-dir',type=Path)
    p.add_argument('--workbook',type=Path)
    args=p.parse_args()
    render(args.game_dir.resolve(),Path(__file__).resolve().parents[1],args.csv_dir,args.workbook)
