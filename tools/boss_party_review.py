#!/usr/bin/env python3
"""Render the pending boss-party review; never writes runtime game data."""
import argparse
import collections
import json
from pathlib import Path

ROLES = {'gym': '관장', 'elite': '사천왕', 'champion': '챔피언'}


def validate(game, data, species):
    def read(name):
        return json.loads((game / 'data' / name).read_text(encoding='utf-8'))
    expansion=read('boss_expansion.json')
    league=read('league.json')
    current={}
    for role, rows in [('gym',read('gym_leaders.json')['leaders']+expansion['gyms']),
                       ('elite',league['elite']+expansion['elite']),
                       ('champion',league['champions']+expansion['champions'])]:
        for row in rows:
            current[role+':'+row.get('leader_id',row.get('id'))]=list(map(int,row['party']))
    entries=data['entries']
    assert data['status']=='review_pending_not_applied'
    assert len(entries)==137 and set(e['key'] for e in entries)==set(current)
    covered=set()
    checks=3
    for e in entries:
        assert e['current_party']==current[e['key']],e['key']
        assert len(e['proposed_party'])==len(current[e['key']]),e['key']
        assert e['proposed_party'][-len(e['fixed_aces']):]==e['fixed_aces'],e['key']
        assert len(e['fixed_aces'])==(2 if e['key']=='gym:tate_liza' else 1),e['key']
        assert e['reference']['sources'] and e['reference']['game'] and e['reference']['battle']
        assert 1<=len(e['reference']['party'])<=6,e['key']
        assert all(url.startswith('https://') for url in e['reference']['sources'])
        for party in [e['current_party'],e['proposed_party'],e['reference']['party'],e['fixed_aces'],e['supplements']]+[a['party'] for a in e['reference']['alternatives']]:
            assert all(type(p) is int and p in species for p in party),e['key']
            checks+=1
        assert set(e['supplements'])<=set(e['current_party']),e['key']
        explained=set(e['reference']['party']+e['fixed_aces']+e['supplements'])
        assert set(e['proposed_party'])<=explained,e['key']
        for change in e['evolution_changes']:
            assert change['from'] in e['reference']['party'] and change['to'] in e['fixed_aces'],e['key']
        covered.update(e['spreadsheet_rows'])
        checks+=9
    for e in data['pending_candidates']:
        assert e['key'] not in current
        covered.update(e['spreadsheet_rows'])
        checks+=1
    assert covered==set(range(4,139))
    assert len(data['spreadsheet_rows'])==135
    checks+=2
    return checks


def render(game, source, destination, workbook=None):
    data = json.loads(source.read_text(encoding='utf-8'))
    dex = json.loads((game / 'data/pokedex_expansion.json').read_text(encoding='utf-8'))
    species = {int(p['id']): p for p in dex['pokemon']}
    checks=validate(game,data,species)

    def names(ids):
        return ' · '.join(species[id]['name'] for id in ids)

    rows = data['entries']
    counts = collections.Counter(e['role'] for e in rows)
    lines = [
        '# 원작 기반 트레이너 파티 검토안', '',
        '**상태: 컨펌 대기. 게임에는 아직 적용하지 않았습니다.**', '',
        f"확인 {data['checked_on']}. 관장{counts['gym']}·사천왕{counts['elite']}·챔피언{counts['champion']}, 총{len(rows)}대진. "
        '첨부 엑셀135행 전체를 대조했습니다. 진우는 아래 추가 후보로 분리했습니다.', '',
        '‘원작 파티’는 표시된 작품/전투에서 실제 사용한 구성이며, ‘게임 검토안’은 엑셀 대표를 마지막에 두고 현재 정원을 유지한 안입니다. '
        '원작에 없는 멤버를 원작 파티처럼 표시하지 않습니다. 굵은 글씨는 고정 대표, `보충`은 원작 사용을 확인하지 않은 기존 게임 멤버입니다.', '',
        '원작에 2~5마리만 있는 경우 임의로 원작 6마리 파티를 만들지 않았습니다. 부족한 자리는 보충안으로 따로 표시했으므로, '
        '원작 인원 그대로 적용하거나 보충안을 바꾸려면 이 표에서 검토할 수 있습니다. 나머지 멤버를 무작위로 뽑는 기능은 추가하지 않았습니다.', '',
        '대표의 미진화체만 엑셀 최종형으로 바꿉니다. 그 밖의 원작 종/상시 폼은 유지하며 메가진화·거다이맥스·테라스탈 전투 효과는 적용하지 않습니다. '
        '장식·색상 폼과 원작의 레벨/기술/도구/특성까지 그대로 이식하는 목록은 아닙니다. 게임 레벨·전술 변경도 아직 하지 않았습니다.', '',
        '[검토 엑셀](boss-party-review.xlsx) · [상세 JSON](boss-party-review.json) · [현재 게임의 역할/BGM 목록](boss-roster.md)', '',
        '## 확인이 필요한 예외', '',
        '- 원작 최종진화와 엑셀 대표를 구분했습니다. 레드는 원작 피카츄를 쓰지만 검토안은 엑셀대로 라이츄입니다.',
        '- 코르니의 체육관 팀에는 루카리오가 없습니다. 별도 마스터타워 이벤트의 루카리오를 고정하는 것은 엑셀 요청에 따른 예외입니다.',
        '- 풍·란은 현재 공동 대진이므로 솔록·루나톤 두 자리를 확보했습니다. 새 대진으로 분리하지 않았습니다.',
        '- 시바는 엑셀에 관동·성도 두 행이 있지만 게임에는 사천왕 역할 하나입니다. 두 행을 같은 역할에 연결했고 지역별 분리는 하지 않았습니다.',
        '- 비트·마리·타임은 엑셀에 없어 원작 대표를 제안했습니다. 진우는 기존 게임에 없어 추가 후보로만 기록했습니다.',
        '- 스타팅에 따라 달라지는 그린·단델·네모·쿠쿠이·하우, 태세가 갈리는 마스터드의 원작 대안은 별도 표에 기록했습니다. 자동 연동은 미구현입니다.', '',
    ]
    regions = list(dict.fromkeys(e['region'] for e in rows))
    for region in regions:
        lines += ['## ' + region, '',
                  '| 역할·이름 | 기준 작품 / 전투 | 원작 파티 | 게임 검토안 | 원작 외 보충 / 차이 |',
                  '|---|---|---|---|---|']
        for e in rows:
            if e['region'] != region:
                continue
            ref = e['reference']
            team = ' · '.join(('**' + species[p]['name'] + '**') if i >= len(e['proposed_party'])-len(e['fixed_aces'])
                              else species[p]['name'] for i, p in enumerate(e['proposed_party']))
            notes = []
            if e['supplements']:
                notes.append('보충: ' + names(e['supplements']))
            for change in e['evolution_changes']:
                notes.append(names([change['from']]) + '→' + names([change['to']]))
            if e['omitted_reference_members']:
                notes.append('정원 조정으로 제외: ' + names(e['omitted_reference_members']))
            if any(p not in ref['party'] and p not in [c['to'] for c in e['evolution_changes']] for p in e['fixed_aces']):
                notes.append('엑셀 대표 별도 추가')
            if not e['spreadsheet_rows']:
                notes.append('엑셀 미수록 · 원작 대표 제안')
            label = f"{ref['game']} / {ref['battle']}"
            lines.append('| ' + ' | '.join([ROLES[e['role']]+' '+e['name'], f"[{label}]({ref['sources'][0]})",
                                           names(ref['party']), team, '<br>'.join(notes) or '대표 위치만 조정'])+' |')
        lines.append('')
    lines += ['## 원작 분기와 역할별 주의사항', '', '| 역할 키 | 주의사항 / 원작 대안 |', '|---|---|']
    for e in rows:
        notes = list(e['notes'])
        notes += [alt['label']+': '+names(alt['party']) for alt in e['reference']['alternatives']]
        if notes:
            lines.append('| `'+e['key']+'` | '+'<br>'.join(notes)+' |')
    lines += ['', '## 기존 게임에 없는 추가 후보', '']
    for e in data['pending_candidates']:
        lines += [f"- **{e['name']}**: {names(e['reference']['party'])}. 고정 대표 {names(e['fixed_aces'])}. {e['status']}.",
                  f"  원작: [{e['reference']['game']}]({e['reference']['sources'][0]})."]
    lines += ['', '## 변경 전 구성과 대조', '', '| 역할 키 | 현재 게임 파티 |', '|---|---|']
    for e in rows:
        lines.append('| `'+e['key']+'` | '+names(e['current_party'])+' |')
    destination.write_text('\n'.join(lines)+'\n', encoding='utf-8')

    if workbook:
        # Optional review artifact only; openpyxl is not a game dependency.
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        wb = Workbook()
        ws = wb.active
        ws.title = '게임 적용 검토안'
        ws.append(['지방', '역할', '이름', '원작 기준', '원작 파티', '고정 대표', '게임 검토안', '원작 외 보충', '현재 게임 파티', '주의사항', '출처', '엑셀 행'])
        for e in rows:
            ref = e['reference']
            ws.append([e['region'], ROLES[e['role']], e['name'], ref['game']+' / '+ref['battle'], names(ref['party']),
                       names(e['fixed_aces']), names(e['proposed_party']), names(e['supplements']), names(e['current_party']),
                       '\n'.join(e['notes']), '\n'.join(ref['sources']), ', '.join(map(str,e['spreadsheet_rows']))])
        ws.freeze_panes = 'D2'
        widths = [14,12,16,36,64,32,70,40,64,65,65,15]
        for i, width in enumerate(widths,1):
            ws.column_dimensions[ws.cell(1,i).column_letter].width=width
        for cell in ws[1]:
            cell.fill = PatternFill('solid',fgColor='19383D')
            cell.font = Font(color='FFFFFF',bold=True)
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=84
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
            row[5].font=Font(bold=True,color='1E6657')
            if row[7].value:
                row[7].fill=PatternFill('solid',fgColor='FFF0CE')
        table=Table(displayName='PartyReview',ref=f'A1:L{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
        info=wb.create_sheet('검토 기준')
        info.append(['항목','내용'])
        info.append(['상태','컨펌 대기 / 게임 미적용 / EXE 미생성'])
        for key,value in data['policy'].items():info.append([key,value])
        info.append(['주의','원작 외 보충 열은 원작 사용 포켓몬으로 확인한 것이 아닙니다. 대표는 마지막, 풍·란만 두 자리입니다.'])
        info.append(['장식 폼','장식·색상 폼 전체를 원작과 대조한 목록이 아닙니다. 명시된 지역/상시 폼을 구분합니다.'])
        for e in data['pending_candidates']:
            info.append(['추가 후보 '+e['name'],e['status']+' / '+names(e['reference']['party'])])
        info.column_dimensions['A'].width=26
        info.column_dimensions['B'].width=110
        for row in info:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        info.freeze_panes='B2'
        alt=wb.create_sheet('원작 분기')
        alt.append(['역할 키','원작 분기','파티','출처'])
        for e in rows+data['pending_candidates']:
            for option in e['reference']['alternatives']:
                alt.append([e['key'],option['label'],names(option['party']),e['reference']['sources'][0]])
        for column,width in [('A',28),('B',45),('C',110),('D',80)]:alt.column_dimensions[column].width=width
        alt.freeze_panes='C2'
        workbook.parent.mkdir(parents=True,exist_ok=True)
        wb.save(workbook)
        from openpyxl import load_workbook
        verified=load_workbook(workbook,read_only=True,data_only=True)
        assert verified['게임 적용 검토안'].max_row==138
        assert verified['게임 적용 검토안'].max_column==12
        verified.close()
        checks+=2
    print(f'Party review: {checks} checks / 0 failures; runtime game data untouched')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--game-dir',required=True,type=Path)
    parser.add_argument('--input',type=Path,default=Path(__file__).resolve().parents[1]/'docs/boss-party-review.json')
    parser.add_argument('--output',type=Path,default=Path(__file__).resolve().parents[1]/'docs/boss-party-review.md')
    parser.add_argument('--workbook',type=Path)
    args=parser.parse_args()
    render(args.game_dir,args.input,args.output,args.workbook)
